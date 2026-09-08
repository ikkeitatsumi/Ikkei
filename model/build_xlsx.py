# -*- coding: utf-8 -*-
"""
収支試算 Excel（数式入り）を生成する。
  python3 model/build_xlsx.py  →  output/西麻布4丁目_収支試算.xlsx

「前提条件」シートの青字セルを変更すると「分譲」「ホテル」シートが再計算される。
「用途比較」「感度分析」「残余法」は feasibility.py の計算結果（値）を転記。
"""
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import feasibility as fz

OUT = os.path.join(os.path.dirname(__file__), "..", "output", "西麻布4丁目_収支試算.xlsx")

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
HDR = PatternFill("solid", fgColor="DDEBF7")
SUB = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN = '#,##0'
OKU = '#,##0.00"億"'
PCT = '0.0%'


def title(ws, text, row=1):
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, size=13)


def hdr(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l)
        c.font = BOLD
        c.fill = HDR
        c.border = BOX
        c.alignment = Alignment(horizontal="center")


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build():
    site, cc, results, sens, far_cases, residual, prices, unit_prices = fz.run(verbose=False)
    wb = Workbook()

    # ------------------------------------------------------------------ 前提条件
    ws = wb.active
    ws.title = "前提条件"
    title(ws, "西麻布4丁目 土地（FACE西麻布） 収支試算 前提条件　※青字は変更可")
    widths(ws, [34, 18, 44])
    hdr(ws, 3, ["項目", "値", "備考"])
    inputs = [
        ("土地面積(㎡)", site.land_m2, "登記簿。実測なし → 要確定測量", YEN),
        ("土地価格(円)", site.asking_price, "募集価格25億円(税込)。価格交渉可", YEN),
        ("近隣商業地域の面積割合", site.share_kinsho, "用途地域図で要確認。一住(60/300)+近商(80/400)の跨り", PCT),
        ("容積率 一住", site.far_ichiju, "", PCT),
        ("容積率 近商", site.far_kinsho, "", PCT),
        ("前面道路幅員(m)", site.front_road_w, "西側公道7.1m。南側私道4mは容積計算に不使用", '0.0'),
        ("道路幅員係数 一住", site.road_coef_ichiju, "住居系 0.4", '0.0'),
        ("道路幅員係数 近商", site.road_coef_kinsho, "その他 0.6", '0.0'),
        ("既存建物延床(㎡)", site.existing_gfa_m2, "RC B1/3F 2002年", YEN),
        ("現行年間賃料(円)", site.existing_rent_year, "レントロール 7,975,000円/月", YEN),
        ("取得〜退去 月数", site.rent_months_until_vacate, "2026/11取得→2027/1末退去想定", '0'),
        ("仲介手数料率", cc.broker_fee_rate, "税抜。別途消費税", PCT),
        ("土地 固定資産税評価(円)", cc.land_tax_value, "公示地価300万/㎡×0.7 想定", YEN),
        ("登録免許税率(土地)", cc.reg_tax_rate, "", PCT),
        ("不動産取得税率(宅地特例後)", cc.acq_tax_rate, "3%×1/2", PCT),
        ("その他取得費用(円)", cc.other_acq_cost, "司法書士・印紙・測量・調査", YEN),
        ("固都税 年額(円)", cc.property_tax_year, "土地+既存建物", YEN),
        ("解体 坪単価(円)", cc.demolition_per_tsubo, "RC・地下あり 東京", YEN),
        ("解体 追加費用(円)", cc.demolition_extra, "アスベスト・埋戻し・近隣対策", YEN),
        ("設計監理率", cc.design_rate, "本体工事比", PCT),
        ("外構・調査・近隣対策等(円)", cc.site_misc, "", YEN),
        ("予備費率", cc.contingency_rate, "本体工事比", PCT),
        ("金利(年)", cc.loan_rate, "", PCT),
        ("融資手数料率", cc.loan_fee, "", PCT),
    ]
    names = {}
    r = 4
    for label, val, note, fmt in inputs:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=val)
        c.font = BLUE
        c.number_format = fmt
        c.border = BOX
        ws.cell(row=r, column=3, value=note).border = BOX
        names[label] = f"'前提条件'!$B${r}"
        r += 1
    # 計算値
    r += 1
    ws.cell(row=r, column=1, value="計算値").font = BOLD
    r += 1
    calc = [
        ("土地面積(坪)", f"={names['土地面積(㎡)']}/3.30579", '0.00'),
        ("土地 坪単価(万円)", f"={names['土地価格(円)']}/B{r}/10000", YEN),
        ("許容容積率(道路幅員反映)",
         f"=(1-{names['近隣商業地域の面積割合']})*MIN({names['容積率 一住']},{names['前面道路幅員(m)']}*{names['道路幅員係数 一住']})"
         f"+{names['近隣商業地域の面積割合']}*MIN({names['容積率 近商']},{names['前面道路幅員(m)']}*{names['道路幅員係数 近商']})", PCT),
        ("容積対象床面積(㎡)", f"={names['土地面積(㎡)']}*B{r + 2}", YEN),
        ("既存建物解体費(円)", f"={names['既存建物延床(㎡)']}/3.30579*{names['解体 坪単価(円)']}+{names['解体 追加費用(円)']}", YEN),
        ("取得諸経費(円)",
         f"={names['土地価格(円)']}*{names['仲介手数料率']}*1.1+{names['土地 固定資産税評価(円)']}*({names['登録免許税率(土地)']}+{names['不動産取得税率(宅地特例後)']})+{names['その他取得費用(円)']}", YEN),
        ("退去までの賃料収入(円)", f"={names['現行年間賃料(円)']}/12*{names['取得〜退去 月数']}", YEN),
    ]
    for label, f, fmt in calc:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=f)
        c.number_format = fmt
        c.border = BOX
        names[label] = f"'前提条件'!$B${r}"
        r += 1

    # ------------------------------------------------------------------ 分譲
    ws = wb.create_sheet("分譲マンション")
    title(ws, "分譲マンション 収支（数式）　※青字は変更可")
    widths(ws, [30, 18, 40])
    hdr(ws, 3, ["前提", "値", "備考"])
    ca = fz.CondoAssume()
    cin = [
        ("延床/容積対象 倍率", ca.gfa_factor, "共用廊下・階段・EV・地下住宅1/3 等の不算入分", '0.00'),
        ("専有/容積対象 倍率", ca.sellable_factor, "", '0.00'),
        ("販売坪単価(円)", ca.price_per_tsubo, "港区2026年1-4月平均758万/坪。西麻布小規模高級で1,500万想定", YEN),
        ("本体工事 坪単価(円)", ca.construction_per_tsubo, "RC造 東京 小規模高級仕様", YEN),
        ("販売経費率(売上比)", ca.sales_cost_rate, "広告・販売手数料", PCT),
        ("事業期間(月)", ca.months, "退去3+解体3+工事18+引渡3", '0'),
    ]
    r = 4
    cn = {}
    for label, val, note, fmt in cin:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=val)
        c.font = BLUE
        c.number_format = fmt
        c.border = BOX
        ws.cell(row=r, column=3, value=note).border = BOX
        cn[label] = f"$B${r}"
        r += 1
    r += 1
    hdr(ws, r, ["項目", "金額(円)", "算式"])
    r += 1
    start = r
    lines = [
        ("土地代", f"={names['土地価格(円)']}"),
        ("取得諸経費", f"={names['取得諸経費(円)']}"),
        ("固都税(事業期間)", f"={names['固都税 年額(円)']}*{cn['事業期間(月)']}/12"),
        ("既存建物解体", f"={names['既存建物解体費(円)']}"),
        ("退去までの賃料収入", f"=-{names['退去までの賃料収入(円)']}"),
        ("本体工事費", f"={names['容積対象床面積(㎡)']}*{cn['延床/容積対象 倍率']}/3.30579*{cn['本体工事 坪単価(円)']}"),
        ("設計監理", f"=B{r + 5}*{names['設計監理率']}"),
        ("外構・調査・近隣対策等", f"={names['外構・調査・近隣対策等(円)']}"),
        ("予備費", f"=B{r + 5}*{names['予備費率']}"),
        ("販売経費", f"=B{r + 16}*{cn['販売経費率(売上比)']}"),
        ("金利・融資手数料",
         f"=(SUM(B{r}:B{r + 4})+SUM(B{r + 6}:B{r + 9}))*{names['金利(年)']}*{cn['事業期間(月)']}/12"
         f"+B{r + 5}*{names['金利(年)']}*{cn['事業期間(月)']}/12*0.5"
         f"+SUM(B{r}:B{r + 9})*{names['融資手数料率']}"),
    ]
    for label, f in lines:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=f)
        c.number_format = YEN
        c.border = BOX
        r += 1
    ws.cell(row=r, column=1, value="総事業費").font = BOLD
    c = ws.cell(row=r, column=2, value=f"=SUM(B{start}:B{r - 1})")
    c.number_format = YEN
    c.font = BOLD
    total_row = r
    r += 2
    hdr(ws, r, ["収入・指標", "値", ""])
    r += 1
    out = [
        ("専有面積(㎡)", f"={names['容積対象床面積(㎡)']}*{cn['専有/容積対象 倍率']}", YEN),
        ("専有面積(坪)", f"=B{r}/3.30579", '0.0'),
        ("売上高(円)", f"=B{r + 1}*{cn['販売坪単価(円)']}", YEN),   # start+16 行
        ("事業利益(円)", f"=B{r + 2}-B{total_row}", YEN),
        ("利益率(売上比)", f"=B{r + 3}/B{r + 2}", PCT),
        ("利益率(総事業費比)", f"=B{r + 3}/B{total_row}", PCT),
        ("損益分岐 販売坪単価(円)", f"=B{total_row}/B{r + 1}", YEN),
        ("想定戸数(約100㎡/戸)", f"=ROUND(B{r}/100,0)", '0'),
    ]
    assert r == start + 14, (r, start)  # 売上高 が start+16 行にあることを保証
    for label, f, fmt in out:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=f)
        c.number_format = fmt
        c.border = BOX
        r += 1

    # ------------------------------------------------------------------ ホテル
    ws = wb.create_sheet("ホテル")
    title(ws, "ホテル 収支（数式）　※青字は変更可。初期値はアップスケール想定")
    widths(ws, [30, 18, 44])
    hdr(ws, 3, ["前提", "値", "備考"])
    ha = fz.HotelAssume()
    hin = [
        ("延床/容積対象 倍率", ha.gfa_factor, "ホテルは廊下も容積算入。機械室のみ不算入", '0.00'),
        ("延床/客室(㎡/室)", ha.gfa_per_key_m2, "宿泊特化28 / アップスケール36 / ラグジュアリー50", '0.0'),
        ("ADR(円)", ha.adr, "宿泊特化28,000 / アップスケール45,000 / ラグジュアリー90,000", YEN),
        ("稼働率", ha.occupancy, "", PCT),
        ("料飲・その他売上比率", ha.other_rev_rate, "客室売上比", PCT),
        ("所有者NOI率(総売上比)", ha.owner_noi_rate, "賃貸借型の賃料 / MC型のGOP−手数料−FF&E積立", PCT),
        ("本体工事 坪単価(円)", ha.construction_per_tsubo, "", YEN),
        ("FF&E(円/室)", ha.ffe_per_key, "", YEN),
        ("開業準備費(円)", ha.preopening, "", YEN),
        ("事業期間(月)", ha.months, "", '0'),
        ("還元利回り(Cap)", ha.exit_cap, "CBRE 2026/3 東京ホテル4.2%に保守性を上乗せ", PCT),
        ("運営時固都税倍率", 1.6, "新築建物分の増加", '0.0'),
    ]
    r = 4
    hn = {}
    for label, val, note, fmt in hin:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=val)
        c.font = BLUE
        c.number_format = fmt
        c.border = BOX
        ws.cell(row=r, column=3, value=note).border = BOX
        hn[label] = f"$B${r}"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="延床面積(㎡)").border = BOX
    c = ws.cell(row=r, column=2, value=f"={names['容積対象床面積(㎡)']}*{hn['延床/容積対象 倍率']}")
    c.number_format = YEN
    gfa_cell = f"$B${r}"
    r += 1
    ws.cell(row=r, column=1, value="客室数").border = BOX
    c = ws.cell(row=r, column=2, value=f"=INT({gfa_cell}/{hn['延床/客室(㎡/室)']})")
    keys_cell = f"$B${r}"
    r += 2
    hdr(ws, r, ["項目", "金額(円)", ""])
    r += 1
    start = r
    lines = [
        ("土地代", f"={names['土地価格(円)']}"),
        ("取得諸経費", f"={names['取得諸経費(円)']}"),
        ("固都税(事業期間)", f"={names['固都税 年額(円)']}*{hn['事業期間(月)']}/12"),
        ("既存建物解体", f"={names['既存建物解体費(円)']}"),
        ("退去までの賃料収入", f"=-{names['退去までの賃料収入(円)']}"),
        ("本体工事費", f"={gfa_cell}/3.30579*{hn['本体工事 坪単価(円)']}"),
        ("設計監理", f"=B{r + 5}*{names['設計監理率']}"),
        ("FF&E", f"={keys_cell}*{hn['FF&E(円/室)']}"),
        ("外構・調査・近隣対策等", f"={names['外構・調査・近隣対策等(円)']}"),
        ("開業準備費", f"={hn['開業準備費(円)']}"),
        ("予備費", f"=B{r + 5}*{names['予備費率']}"),
        ("金利・融資手数料",
         f"=(SUM(B{r}:B{r + 4})+SUM(B{r + 6}:B{r + 10}))*{names['金利(年)']}*{hn['事業期間(月)']}/12"
         f"+B{r + 5}*{names['金利(年)']}*{hn['事業期間(月)']}/12*0.5"
         f"+SUM(B{r}:B{r + 10})*{names['融資手数料率']}"),
    ]
    for label, f in lines:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=f)
        c.number_format = YEN
        c.border = BOX
        r += 1
    ws.cell(row=r, column=1, value="総事業費").font = BOLD
    c = ws.cell(row=r, column=2, value=f"=SUM(B{start}:B{r - 1})")
    c.number_format = YEN
    c.font = BOLD
    total_row = r
    r += 2
    hdr(ws, r, ["収入・指標", "値", ""])
    r += 1
    out = [
        ("RevPAR(円)", f"={hn['ADR(円)']}*{hn['稼働率']}", YEN),
        ("客室売上(円/年)", f"={keys_cell}*{hn['ADR(円)']}*{hn['稼働率']}*365", YEN),
        ("総売上(円/年)", f"=B{r + 1}*(1+{hn['料飲・その他売上比率']})", YEN),
        ("所有者NOI(円/年)", f"=B{r + 2}*{hn['所有者NOI率(総売上比)']}-{names['固都税 年額(円)']}*{hn['運営時固都税倍率']}", YEN),
        ("NOI利回り(総事業費比)", f"=B{r + 3}/B{total_row}", PCT),
        ("収益価格(円)", f"=B{r + 3}/{hn['還元利回り(Cap)']}", YEN),
        ("含み損益(収益価格−総事業費)", f"=B{r + 5}-B{total_row}", YEN),
        ("損益分岐ADR(円)",
         f"=((B{total_row}*{hn['還元利回り(Cap)']}+{names['固都税 年額(円)']}*{hn['運営時固都税倍率']})/{hn['所有者NOI率(総売上比)']}/(1+{hn['料飲・その他売上比率']}))/({keys_cell}*{hn['稼働率']}*365)", YEN),
    ]
    for label, f, fmt in out:
        ws.cell(row=r, column=1, value=label).border = BOX
        c = ws.cell(row=r, column=2, value=f)
        c.number_format = fmt
        c.border = BOX
        r += 1

    # ------------------------------------------------------------------ 用途比較（値）
    ws = wb.create_sheet("用途比較")
    title(ws, "用途別 収支比較（土地25億円・基準ケース、feasibility.py の計算値）")
    widths(ws, [26, 16, 16, 16, 16, 16, 18, 16])
    hdr(ws, 3, ["用途", "総事業費", "年間売上/売上高", "年間NOI/事業利益", "利回り/利益率", "収益価格", "含み損益", "備考"])
    r = 4
    for key, res in results.items():
        m = res.metrics
        cap_key = next((k for k in m if k.startswith("収益価格")), None)
        if res.kind == "sale":
            row = [res.name, res.total_cost, res.revenue, res.profit, m["利益率(売上比)"], None, None, f"専有{m['専有面積坪']:.0f}坪・{m['想定戸数']}戸"]
        else:
            note = ""
            if "客室数" in m:
                note = f"{m['客室数']}室・ADR{m['ADR']:,.0f}円・稼働{m['稼働率'] * 100:.0f}%"
            elif "貸室面積坪" in m:
                note = f"貸室{m['貸室面積坪']:.0f}坪・{m['賃料坪単価/月']:,.0f}円/坪"
            elif "現行年間賃料" in m:
                note = "再募集後ベース"
            row = [res.name, res.total_cost, res.revenue, res.noi, m.get("NOI利回り(総事業費比)", m.get("再募集後NOI利回り")), m[cap_key], m["含み損益"], note]
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            if i in (2, 3, 4, 6, 7) and v is not None:
                c.number_format = OKU
                c.value = v / fz.OKU
            if i == 5 and v is not None:
                c.number_format = PCT
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="※ 保有型は「収益価格＝NOI÷還元利回り」と総事業費の差を含み損益として表示。ホテル4.5%・賃貸住宅3.7%・店舗4.0%・現況保有3.5%。").font = Font(italic=True)

    # ------------------------------------------------------------------ 感度分析（値）
    ws = wb.create_sheet("感度分析")
    title(ws, "分譲マンション 事業利益（億円）と利益率　行：販売坪単価 / 列：土地価格")
    widths(ws, [18] + [16] * len(prices) * 2)
    hdr(ws, 3, ["販売坪単価"] + [f"土地{p / fz.OKU:.0f}億 利益" for p in prices] + [f"土地{p / fz.OKU:.0f}億 利益率" for p in prices])
    r = 4
    for up, row in sens:
        ws.cell(row=r, column=1, value=f"{up / fz.MAN:,.0f}万円/坪").border = BOX
        for j, (v, m) in enumerate(row):
            c = ws.cell(row=r, column=2 + j, value=v / fz.OKU)
            c.number_format = '0.0'
            c.border = BOX
            c2 = ws.cell(row=r, column=2 + len(prices) + j, value=m)
            c2.number_format = PCT
            c2.border = BOX
            if m < 0.10:
                c.fill = PatternFill("solid", fgColor="F8CBAD")
                c2.fill = PatternFill("solid", fgColor="F8CBAD")
            elif m >= 0.15:
                c.fill = PatternFill("solid", fgColor="C6EFCE")
                c2.fill = PatternFill("solid", fgColor="C6EFCE")
        r += 1
    r += 2
    ws.cell(row=r, column=1, value="容積率ケース（分譲・坪1,500万・土地25億）").font = BOLD
    r += 1
    hdr(ws, r, ["ケース", "許容容積率", "専有面積(坪)", "事業利益(億)", "利益率"])
    r += 1
    for label, far, tsubo, pr, mg in far_cases:
        vals = [label, far, tsubo, pr / fz.OKU, mg]
        fmts = [None, PCT, '0', '0.00', PCT]
        for i, (v, f) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BOX
            if f:
                c.number_format = f
        r += 1

    # ------------------------------------------------------------------ 残余法（値）
    ws = wb.create_sheet("残余法")
    title(ws, "残余法による土地価格（用途別に支払可能な上限）")
    widths(ws, [40, 16, 18])
    hdr(ws, 3, ["用途・条件", "土地価格(億円)", "坪単価(万円)"])
    r = 4
    for k, v in residual.items():
        ws.cell(row=r, column=1, value=k).border = BOX
        c = ws.cell(row=r, column=2, value=v / fz.OKU)
        c.number_format = '0.00'
        c.border = BOX
        c = ws.cell(row=r, column=3, value=v / site.land_tsubo / fz.MAN)
        c.number_format = YEN
        c.border = BOX
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"募集価格 25.00億円 = {site.asking_price / site.land_tsubo / fz.MAN:,.0f}万円/坪").font = BOLD

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("saved", os.path.abspath(OUT))


if __name__ == "__main__":
    build()
